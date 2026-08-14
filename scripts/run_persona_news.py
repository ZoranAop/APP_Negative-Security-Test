#!/usr/bin/env python3
"""
run_persona_news.py — 按人设生成+发布脚本

从 200 账号表（含 persona_id / persona_name / persona_desc / persona_post_style /
persona_post_lang）中按话题选择人设匹配的用户，采集相关资讯（Google News RSS），
按人设的语言与口吻生成第一人称文案，再通过 publish_from_tokens.py 发布纯文本帖。

用法：
    # 从 Excel 读账号，选 ai_tech/digital 人设发 3 条科技资讯
    py -3 scripts/run_persona_news.py --accounts-xlsx "D:/.../pre_运营用户200昵称.xlsx" \
        --topic tech --personas ai_tech,digital --count 3 --yes

    # 从 CSV 读账号（列：邮箱,用户名,昵称,密码,pincode,persona_id,persona_name,persona_desc,persona_post_style,persona_post_lang）
    py -3 scripts/run_persona_news.py --accounts-csv accounts200.csv \
        --topic finance --personas ai_tech --count 3 --yes
"""
from __future__ import annotations

import argparse
import csv
import html
import io
import re
import subprocess
import sys
import urllib.parse
from pathlib import Path

import requests

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
PY = [sys.executable]

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36")

# ---------------------------------------------------------------------------
# 话题配置：Google News 查询词 + 话题标签 + 话题 → 推荐人设
# ---------------------------------------------------------------------------
TOPICS = {
    "tech": {
        "query": "AI 大模型 芯片 发布",
        "tags": {"zh": "#AI #科技 #大模型", "zh_hant": "#AI #科技 #大模型",
                 "en": "#AI #Tech #LLM", "ja": "#AI #テクノロジー #大規模モデル"},
        "personas": ["ai_tech", "digital"],
    },
    "finance": {
        "query": "财经 股市 黄金 港股",
        "tags": {"zh": "#财经 #股市 #黄金", "zh_hant": "#財經 #股市 #黃金",
                 "en": "#Finance #Market #Gold", "ja": "#経済 #株 #金"},
        "personas": ["ai_tech", "digital"],
    },
    "entertainment": {
        "query": "娱乐 明星 演唱会 票房",
        "tags": {"zh": "#娱乐 #演唱会 #明星", "zh_hant": "#娛樂 #演唱會 #明星",
                 "en": "#Entertainment #Concert #Celebrity", "ja": "#エンタメ #ライブ #芸能"},
        "personas": ["music", "beauty", "fashion"],
    },
    "game": {
        "query": "游戏 手游 二次元 新版本",
        "tags": {"zh": "#游戏 #手游 #二次元", "zh_hant": "#遊戲 #手遊 #二次元",
                 "en": "#Gaming #MobileGame", "ja": "#ゲーム #アプリ"},
        "personas": ["gaming"],
    },
    "science": {
        "query": "科学 天文 自然 发现",
        "tags": {"zh": "#科学 #天文 #科普", "zh_hant": "#科學 #天文 #科普",
                 "en": "#Science #Astronomy", "ja": "#科学 #宇宙"},
        "personas": ["science"],
    },
    "sports": {
        "query": "体育 足球 篮球 赛事",
        "tags": {"zh": "#体育 #足球 #篮球", "zh_hant": "#體育 #足球 #籃球",
                 "en": "#Sports #Football", "ja": "#スポーツ #サッカー"},
        "personas": ["fitness"],
    },
}

# 各语言「发帖者口吻」评论模板（第一人称，简短，配合资讯标题）
LANG_COMMENTS = {
    "zh": [
        "这个真的有点东西，值得关注",
        "刷到这条忍不住想说两句",
        "最近这个话题好火，我也来凑个热闹",
        "看完只能说一句：卷，真的卷",
        "随手转一下，懂的都懂",
    ],
    "zh_hant": [
        "這個真的有點東西，值得關注",
        "刷到這條忍不住想說兩句",
        "最近這個話題好火，我也來湊個熱鬧",
        "看完只能說一句：捲，真的捲",
        "隨手轉一下，懂的都懂",
    ],
    "en": [
        "This one's actually interesting, worth a look",
        "Had to comment on this",
        "This topic is everywhere right now",
        "All I can say is it's getting intense",
        "Sharing this for the curious ones",
    ],
    "ja": [
        "これはちょっと気になる",
        "思わずコメントしたくなった",
        "最近この話題、すごく盛り上がってる",
        "見てて一言：すごいことになってる",
        "とりあえずシェアしておく",
    ],
}


def _u8():
    try:
        sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", write_through=True)
        sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", write_through=True)
    except Exception:  # noqa: BLE001
        pass


# ---------------------------------------------------------------------------
# 账号加载（xlsx / csv）
# ---------------------------------------------------------------------------
def load_accounts(path: str) -> list[dict]:
    p = Path(path)
    if p.suffix.lower() == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(p, read_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        header = [str(h or "").strip() for h in rows[0]]
        out = []
        for r in rows[1:]:
            d = {header[i]: (r[i] if i < len(r) else "") for i in range(len(header))}
            out.append(d)
        return out
    with open(p, encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def pick_users(accounts: list[dict], personas: list[str], count: int) -> list[dict]:
    pool = [a for a in accounts if (a.get("persona_id") or "").strip() in personas]
    if not pool:
        pool = accounts
    import random
    if count > 0 and len(pool) > count:
        pool = random.sample(pool, count)
    return pool


# ---------------------------------------------------------------------------
# 资讯采集（Google News RSS）
# ---------------------------------------------------------------------------
def clean_title(t: str) -> str:
    t = html.unescape(t or "")
    t = re.sub(r"<[^>]+>", "", t)
    t = re.sub(r"\s*-\s*[^-]+$", "", t)  # 去 " - 来源"
    return re.sub(r"\s+", " ", t).strip()


def fetch_news(topic: str, want: int) -> list[str]:
    q = TOPICS[topic]["query"]
    url = ("https://news.google.com/rss/search?q=" + urllib.parse.quote(q) +
           "&hl=zh-CN&gl=CN&ceid=CN:zh-Hans")
    r = requests.get(url, headers={"User-Agent": UA}, timeout=25)
    r.raise_for_status()
    txt = r.content.decode("utf-8", errors="ignore")
    titles, seen = [], set()
    for it in re.findall(r"<item>(.*?)</item>", txt, re.S):
        m = re.search(r"<title>(?:<!\[CDATA\[)?(.*?)(?:\]\]>)?</title>", it, re.S)
        t = clean_title(m.group(1)) if m else ""
        if len(t) < 6 or t in seen:
            continue
        seen.add(t)
        titles.append(t)
        if len(titles) >= want:
            break
    return titles


# ---------------------------------------------------------------------------
# 文案生成
# ---------------------------------------------------------------------------
def llm_caption(title: str, account: dict, topic: str) -> tuple[str, str] | None:
    """用 LLM 按人设生成文案（需配置 LLM_TEXT_* / LLM_* 环境变量）。"""
    import os
    api_base = os.getenv("LLM_TEXT_API_BASE") or os.getenv("LLM_API_BASE")
    api_key = os.getenv("LLM_TEXT_API_KEY") or os.getenv("LLM_API_KEY")
    model = os.getenv("LLM_TEXT_MODEL") or os.getenv("LLM_MODEL")
    if not (api_base and api_key and model) or "your_" in api_key:
        return None
    lang = (account.get("persona_post_lang") or "zh").strip()
    persona = account.get("persona_name") or ""
    desc = account.get("persona_desc") or ""
    style = account.get("persona_post_style") or ""
    prompt = (
        f"你是社交媒体上的「{persona}」（关注：{desc}，发帖风格：{style}）。"
        f"请围绕下面这条资讯，用{lang}写一条简短的第一人称分享帖（≤80字），"
        f"口语化、有个人态度，结尾带 3-5 个话题标签。不要提任何媒体名。\n资讯：{title}"
    )
    try:
        r = requests.post(
            api_base.rstrip("/") + "/chat/completions",
            headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
            json={"model": model, "messages": [{"role": "user", "content": prompt}],
                  "temperature": 0.9, "max_tokens": 200},
            timeout=40)
        r.raise_for_status()
        text = r.json()["choices"][0]["message"]["content"].strip()
        return (text, lang) if text else None
    except Exception:  # noqa: BLE001
        return None


def gen_caption(title: str, account: dict, topic: str, use_llm: bool = False) -> tuple[str, str]:
    lang = (account.get("persona_post_lang") or "zh").strip()
    if lang not in LANG_COMMENTS:
        lang = "zh"
    if use_llm:
        got = llm_caption(title, account, topic)
        if got:
            return got
    import random
    comment = random.choice(LANG_COMMENTS[lang])
    tags = TOPICS[topic]["tags"].get(lang, TOPICS[topic]["tags"]["zh"])
    # 中文站标题统一转繁体，避免简繁混搭
    if lang == "zh_hant":
        try:
            import zhconv
            title = zhconv.convert(title, "zh-hant")
        except Exception:  # noqa: BLE001
            pass
    caption = f"{title}，{comment} {tags}".replace("  ", " ").strip()
    return caption, lang


def _run(cmd, env_extra=None):
    import os
    env = os.environ.copy()
    env.setdefault("PYTHONIOENCODING", "utf-8")
    if env_extra:
        env.update(env_extra)
    print("  $ " + " ".join(str(c) for c in cmd), flush=True)
    return subprocess.call(cmd, cwd=str(ROOT), env=env)


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------
def main() -> int:
    _u8()
    ap = argparse.ArgumentParser(description="按人设生成+发布（话题资讯 → 人设文案 → 发布）")
    ap.add_argument("--accounts-xlsx", help="200 账号 Excel 路径")
    ap.add_argument("--accounts-csv", help="200 账号 CSV 路径")
    ap.add_argument("--topic", default="tech", choices=list(TOPICS.keys()))
    ap.add_argument("--personas", default="", help="人设 ID，逗号分隔；留空则用话题默认人设")
    ap.add_argument("--count", type=int, default=3, help="发帖数（= 选用户数）")
    ap.add_argument("--concurrency", type=int, default=3)
    ap.add_argument("--login-spacing", type=float, default=2.5)
    ap.add_argument("--tokens", default="result/tokens_persona.json")
    ap.add_argument("--workdir", default="persona_run")
    ap.add_argument("--use-llm", action="store_true", help="用 LLM 按人设生成文案（需配置 LLM_TEXT_*/LLM_*）")
    ap.add_argument("--skip-publish", action="store_true")
    ap.add_argument("--yes", action="store_true")
    args = ap.parse_args()

    if not args.accounts_xlsx and not args.accounts_csv:
        print("[persona] 需要 --accounts-xlsx 或 --accounts-csv", file=sys.stderr)
        return 2

    accounts = load_accounts(args.accounts_xlsx or args.accounts_csv)
    print(f"[persona] 加载 {len(accounts)} 个账号")

    topic = args.topic
    personas = [p.strip() for p in args.personas.split(",") if p.strip()]
    if not personas:
        personas = TOPICS[topic]["personas"]
    users = pick_users(accounts, personas, args.count)
    if not users:
        print("[persona] 未找到匹配人设的用户", file=sys.stderr)
        return 2
    print(f"[persona] 话题={topic} 人设={personas} 用户 {len(users)} 个")

    # 采集
    titles = fetch_news(topic, len(users))
    if len(titles) < len(users):
        print(f"[persona] 资讯不足: 需要 {len(users)}, 实际 {len(titles)}", file=sys.stderr)
    print(f"[persona] 采集 {len(titles)} 条资讯")

    # 生成文案
    import time
    ts = time.strftime("%Y%m%d_%H%M%S")
    wd = ROOT / args.workdir
    wd.mkdir(parents=True, exist_ok=True)

    acc_csv = wd / f"accounts_{topic}_{ts}.csv"
    moments_csv = wd / f"moments_{topic}_{ts}.csv"

    seen_caps = set()
    rows = []
    for i, u in enumerate(users):
        title = titles[i] if i < len(titles) else titles[-1]
        cap, lang = gen_caption(title, u, topic, use_llm=args.use_llm)
        attempts = 0
        while cap in seen_caps and attempts < 6:
            cap, lang = gen_caption(title, u, topic, use_llm=args.use_llm)
            attempts += 1
        seen_caps.add(cap)
        rows.append({"content": cap, "visibility": 0, "room_id": "", "image_urls": "",
                     "location_name": "", "location_address": "",
                     "location_lat": "", "location_lon": "",
                     "_lang": lang, "_tag": topic})
        print(f"  #{i+1} [{(u.get('昵称') or u.get('邮箱'))[:16]} /{lang}/] {cap[:56]}")

    # 账号 CSV（保留 邮箱/密码/昵称）
    with acc_csv.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["邮箱", "用户名", "昵称", "密码"])
        w.writeheader()
        for u in users:
            w.writerow({"邮箱": u.get("邮箱", ""), "用户名": u.get("用户名", ""),
                        "昵称": u.get("昵称", ""), "密码": u.get("密码", "")})

    fields = ["content", "visibility", "room_id", "image_urls",
              "location_name", "location_address", "location_lat", "location_lon",
              "_lang", "_tag"]
    with moments_csv.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in rows:
            w.writerow(r)
    print(f"[persona] moments 就绪：{len(rows)} 条 → {moments_csv}")

    if args.skip_publish:
        print("[persona] --skip-publish，仅产出素材")
        return 0

    if not args.yes:
        if input("  确认发布真实帖子？输入 y 继续：").strip().lower() not in ("y", "yes"):
            print("[persona] 已取消")
            return 0

    cmd = PY + [str(HERE / "publish_from_tokens.py"),
                "--accounts-csv", str(acc_csv), "--csv", str(moments_csv),
                "--concurrency", str(args.concurrency),
                "--login-spacing", str(args.login_spacing),
                "--tokens-in", str(ROOT / args.tokens),
                "--tokens-out", str(ROOT / args.tokens)]
    rc = _run(cmd, env_extra={"POST_CROP_BOTTOM_HOSTS": ""})
    if rc != 0:
        print("[persona] 发布返回非零", file=sys.stderr)
        return rc
    print("\n[persona] 完成。发布报告见 result/publish_*.csv")
    return 0


if __name__ == "__main__":
    sys.exit(main())
