#!/usr/bin/env python3
"""
run_comments_regional_2b.py — Add comments to the 12 batch-2 regional-news
posts (result/publish_20260915_113730.csv) using FRESH non-publisher token
holders. 3-15 comments/post, random 6-language mix, concise commenter voice.
"""
from __future__ import annotations

import csv, json, random, subprocess, sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
PY = [sys.executable]

REPORT = ROOT / "result" / "publish_20260915_113730.csv"
TOKENS = ROOT / "result" / "tokens.json"
INTERACT_XLSX = ROOT / "互动用户池_670账号.xlsx"
POST_COMMENTS = HERE / "post_comments.py"
BATCH_OUT = ROOT / "temp_comments_regional_2b_batch.csv"

# Concise, highlight-first comment banks by topic (en / zh_hant / ja / ko / ms / id)
BANKS = {
    "politics": {
        "en": [
            "That's the headline to watch right now.",
            "Strong read on the policy direction.",
            "This will define the next cycle, agreed.",
            "The numbers back this up.",
            "Following this closely.",
        ],
        "ja": ["これは注目すべきポイントですね。", "政策の方向性、的確な読みだと思います。", "次のサイクルを左右しますね。"],
        "ko": ["지금 가장 주목할 포인트죠.", "정책 방향을 정확히 보고 계시네요.", "다음 사이클을 좌우할 이슈입니다."],
        "ms": ["Ini memang headline paling penting sekarang.", "Bacaan policy direction tu tepat.", "Ini akan tentukan cycle depan."],
        "id": ["Ini memang headline yang harus dipantau.", "Pembacaan arah kebijakan yang tepat.", "Ini akan menentukan siklus berikutnya."],
        "zh_hant": ["這是現在最值得關注的主線。", "對政策方向判斷得很準。"],
    },
    "tech": {
        "en": [
            "Compute supply is the real constraint now.",
            "This buildout pace is genuinely surprising.",
            "The R&D pipeline here is the story.",
            "Underrated trend, this one.",
            "Whoever controls capacity wins the decade.",
        ],
        "ja": ["コンピュート供給が一番の制約ですね。", "この拡張スピードは正直驚き。", "R&Dパイプラインが本当の物語。"],
        "ko": ["컴퓨트 공급이 진짜 병목이죠.", "이 확장 속도는 정말 놀랍습니다.", "R&D 파이프라인이 핵심 스토리입니다."],
        "ms": ["Supply komput memang constraint sebenar.", "Kelajuan buildout ni memang mengejutkan.", "Pipeline R&D lah cerita sebenar."],
        "id": ["Supply kompute memang bottleneck utama.", "Laju buildout-nya benar-benar mengejutkan.", "Pipeline R&D adalah cerita sesungguhnya."],
        "zh_hant": ["算力供給才是現在的瓶頸。", "這個建設速度確實超出預期。"],
    },
    "web3": {
        "en": [
            "Clear licensing changes everything for institutions.",
            "The on-ramp infrastructure is finally here.",
            "This is the shift that actually matters.",
            "Momentum is real and accelerating.",
            "Institutional entry is opening up.",
        ],
        "ja": ["明確なライセンスが機関にとってすべてを変える。", "機関の上陸路がようやく整った。", "ここが本当の転換点。"],
        "ko": ["명확한 라이선싱이 기관에겐 모든 것을 바꿉니다.", "기관 진입 경로가 마침내 열렸죠.", "진짜 전환점이 바로 여기입니다."],
        "ms": ["Licensing yang jelas ubah segalanya untuk institusi.", "On-ramp infra dah ada akhirnya.", "Ini shift yang betul-betul penting."],
        "id": ["Licensing yang jelas mengubah segalanya bagi institusi.", "On-ramp infrastruktur akhirnya ada.", "Ini pergeseran yang benar-benar penting."],
        "zh_hant": ["明確的牌照對機構端意義重大。", "機構的上陸路終於鋪好。"],
    },
}


def _roll_lang():
    r = random.random()
    if r < 0.5:
        return "en"
    elif r < 0.62:
        return "zh_hant"
    elif r < 0.74:
        return "ja"
    elif r < 0.86:
        return "ko"
    elif r < 0.93:
        return "ms"
    else:
        return "id"


def main() -> int:
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--min-comments", type=int, default=3)
    ap.add_argument("--max-comments", type=int, default=15)
    ap.add_argument("--delay", type=float, default=1.0)
    ap.add_argument("--yes", action="store_true")
    args = ap.parse_args()

    report = list(csv.DictReader(REPORT.open(encoding="utf-8-sig")))
    posts, publishers = [], set()
    for r in report:
        pid = (r.get("moment_id_or_err") or "").strip()
        email = (r.get("email") or "").strip()
        topic = (r.get("_source") or "general").split("_")[-1]
        if pid.isdigit() and "True" in str(r.get("success", "")):
            posts.append({"post_id": int(pid), "topic": topic, "email": email})
            if email:
                publishers.add(email)

    print(f"[posts] {len(posts)} posts, {len(publishers)} publishers to exclude")

    tokens = json.loads(TOKENS.read_text(encoding="utf-8"))
    fresh = [e for e in tokens.keys() if e not in publishers]
    random.shuffle(fresh)
    print(f"[commenters] {len(fresh)} fresh accounts")

    batch = []
    for p in posts:
        n = random.randint(args.min_comments, args.max_comments)
        bank = BANKS.get(p["topic"], BANKS["web3"])
        for _ in range(n):
            lang = _roll_lang()
            text = random.choice(bank.get(lang, bank["en"]))
            batch.append({"email": random.choice(fresh), "post_id": p["post_id"],
                         "topic": p["topic"], "text": text, "nickname": ""})

    random.shuffle(batch)
    with BATCH_OUT.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["email", "post_id", "topic", "text", "nickname"])
        w.writeheader()
        for row in batch:
            w.writerow(row)

    from collections import Counter
    per_post = Counter(r["post_id"] for r in batch)
    print(f"\n[batch] {len(batch)} comments across {len(posts)} posts")
    print(f"  per-post range: {min(per_post.values())}-{max(per_post.values())}")

    if not args.yes:
        if input("  Confirm post comments? (y/n): ").strip().lower() not in ("y", "yes"):
            print("Cancelled."); return 0

    accounts_csv = ROOT / "accounts_regional_2b_commenters.csv"
    import openpyxl
    wb = openpyxl.load_workbook(str(INTERACT_XLSX), read_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    pwd_map = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(headers, row))
        pwd_map[str(d.get("邮箱", "") or "").strip()] = str(d.get("密码", "") or "")
    wb.close()
    with accounts_csv.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["序号", "昵称", "邮箱", "密码"])
        w.writeheader()
        for i, e in enumerate(fresh[:200]):
            w.writerow({"序号": i + 1, "昵称": "", "邮箱": e, "密码": pwd_map.get(e, "")})

    cmd = PY + [str(POST_COMMENTS),
                "--batch", str(BATCH_OUT),
                "--tokens", str(TOKENS),
                "--accounts", str(accounts_csv),
                "--delay", str(args.delay)]
    print(f"\n[exec] {' '.join(str(c) for c in cmd)}")
    rc = subprocess.call(cmd, cwd=str(ROOT))
    return rc


if __name__ == "__main__":
    sys.exit(main())
