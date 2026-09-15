#!/usr/bin/env python3
"""
run_comments_regional_3.py — Add comments to the 15 batch-3 regional-news
posts (result/publish_20260915_115247.csv) using FRESH non-publisher token
holders. 3-15 comments/post, random 8-language mix, concise commenter voice.
"""
from __future__ import annotations

import csv, json, random, subprocess, sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
PY = [sys.executable]

REPORT = ROOT / "result" / "publish_20260915_115247.csv"
TOKENS = ROOT / "result" / "tokens.json"
INTERACT_XLSX = ROOT / "互动用户池_670账号.xlsx"
POST_COMMENTS = HERE / "post_comments.py"
BATCH_OUT = ROOT / "temp_comments_regional_3_batch.csv"

# Concise highlight-first comment banks by topic (8 languages)
BANKS = {
    "politics": {
        "en": [
            "That's the thing that matters most here.",
            "Strong read on the capex shift.",
            "The growth story is getting more durable.",
            "Following this closely — the multiplier is key.",
            "Agreed, this defines the next cycle.",
        ],
        "vi": ["Đây mới là điểm quan trọng nhất.", "Đúng, yếu tố tăng trưởng công nghiệp là then chốt.", "Tôi theo dõi sát hướng này."],
        "pt": ["Isso é o ponto mais importante aqui.", "Uma leitura forte da mudança de capex.", "A história de crescimento fica mais sólida."],
        "ar": ["هذا هو الأهم هنا.", "قراءة قوية لسياسة الرأسمال.", "قصة النمو تصبح أكثر ثباتاً."],
        "ja": ["ここが最も重要ですね。", "政策判断は的確だと思います。"],
        "ko": ["여기가 가장 핵심이죠.", "정책 방향을 정확히 보고 계시네요."],
        "ms": ["Ini memang yang paling penting.", "Bacaan capex tu tepat."],
        "id": ["Ini memang hal terpenting.", "Pembacaan arah capex tepat."],
    },
    "tech": {
        "en": [
            "The data-center + renewable combo is a real differentiator.",
            "Talent retention is the real question here.",
            "Hyperscaler money is aligning with local capability.",
            "This is where the next decade's leverage is.",
            "Underrated market, honestly.",
        ],
        "vi": ["Kết hợp dữ liệu và năng lượng tái tạo là lợi thế thật sự.", "Bảo toàn nhân tài là câu hỏi then chốt.", "Tiền của hyperscaler đang align với năng lực địa phương."],
        "pt": ["A combo de data center e renovável é um diferencial real.", "Retenção de talento é a verdadeira questão.", "O dinheiro dos hyperscalers está alinhando com a capacidade local."],
        "ar": ["دمج مراكز البيانات والطاقة المتجددة ميزة حقيقية.", "الاستثمار هنا هو رافعة العقد القادم."],
        "ja": ["データセンター＋再生エネルギーの掛け合わせが強み。", "人材定着が本当の課題。"],
        "ko": ["데이터센터와 재생에너지 조합이 진짜 차별화 포인트입니다.", "인력 유지가 핵심 과제죠."],
        "ms": ["Combo data center + renewable memang differentiator sebenar.", "Retensi talenta memang soalan utama."],
        "id": ["Kombinasi data center dan energi terbarukan adalah pembeda nyata.", "Retensi talenta memang tantangan utama."],
    },
    "web3": {
        "en": [
            "Clear licensing opens the whole institutional door.",
            "The retail scene is loud — that's the upside.",
            "Compliance path clarity is the unlock.",
            "Institutional entry is finally becoming real.",
            "The momentum is worth watching.",
        ],
        "vi": ["Licensing rõ ràng mở cánh cửa cho tổ chức.", "Dòng vốn trong nước sẽ có lối vào chính thống.", "Momentum đáng theo dõi."],
        "pt": ["Licensa clara abre a porta institucional toda.", "O caminho de compliance é o desbloqueio.", "A entrada institucional finalmente fica real."],
        "ar": ["التراخيص الميسّرة تفتح الباب المؤسساتي.", "وضوح مسار الامتثال هو المفاتيح.", "الزخم يستحق المتابعة."],
        "ja": ["明確なライセンスが機関の入口を開く。", "コンプライアンスの明確化が鍵。"],
        "ko": ["명확한 라이선싱이 기관의 문을 엽니다.", "컴플라이언스 경로 명확화가 핵심입니다."],
        "ms": ["Licensing yang jelas buka pintu institusi.", "Keterang compliance adalah unlock dia.", "Momentum tu memang worth tonton."],
        "id": ["Licensing yang jelas membuka pintu institusi.", "Keterang compliance adalah kuncinya.", "Momentum-nya layak dipantau."],
    },
}


def _roll_lang():
    r = random.random()
    if r < 0.45:
        return "en"
    elif r < 0.55:
        return "vi"
    elif r < 0.65:
        return "pt"
    elif r < 0.73:
        return "ar"
    elif r < 0.81:
        return "ja"
    elif r < 0.89:
        return "ko"
    elif r < 0.95:
        return "ms"
    else:
        return "id"


def main() -> int:
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--min-comments", type=int, default=3)
    ap.add_argument("--max-comments", type=int, default=15)
    ap.add_argument("--delay", type=float, default=1.0)
    ap.add_argument("--yes", action="store_true")
    args = ap.parse_args()

    report = list(csv.DictReader(REPORT.open(encoding="utf-8-sig")))
    posts, publishers = [], set()
    for r in report:
        pid = (r.get("moment_id_or_err") or "").strip()
        email = (r.get("email") or "").strip()
        topic = (r.get("_source") or "general").split("_")[-1]
        if pid.isdigit() and "True" in str(r.get("success", "")):
            posts.append({"post_id": int(pid), "topic": topic, "email": email})
            if email:
                publishers.add(email)

    print(f"[posts] {len(posts)} posts, {len(publishers)} publishers to exclude")

    tokens = json.loads(TOKENS.read_text(encoding="utf-8"))
    fresh = [e for e in tokens.keys() if e not in publishers]
    random.shuffle(fresh)
    print(f"[commenters] {len(fresh)} fresh accounts")

    batch = []
    for p in posts:
        n = random.randint(args.min_comments, args.max_comments)
        bank = BANKS.get(p["topic"], BANKS["web3"])
        for _ in range(n):
            lang = _roll_lang()
            text = random.choice(bank.get(lang, bank["en"]))
            batch.append({"email": random.choice(fresh), "post_id": p["post_id"],
                         "topic": p["topic"], "text": text, "nickname": ""})

    random.shuffle(batch)
    with BATCH_OUT.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["email", "post_id", "topic", "text", "nickname"])
        w.writeheader()
        for row in batch:
            w.writerow(row)

    from collections import Counter
    per_post = Counter(r["post_id"] for r in batch)
    print(f"\n[batch] {len(batch)} comments across {len(posts)} posts")
    print(f"  per-post range: {min(per_post.values())}-{max(per_post.values())}")

    if not args.yes:
        if input("  Confirm post comments? (y/n): ").strip().lower() not in ("y", "yes"):
            print("Cancelled."); return 0

    accounts_csv = ROOT / "accounts_regional_3_commenters.csv"
    import openpyxl
    wb = openpyxl.load_workbook(str(INTERACT_XLSX), read_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    pwd_map = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(headers, row))
        pwd_map[str(d.get("邮箱", "") or "").strip()] = str(d.get("密码", "") or "")
    wb.close()
    with accounts_csv.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["序号", "昵称", "邮箱", "密码"])
        w.writeheader()
        for i, e in enumerate(fresh[:200]):
            w.writerow({"序号": i + 1, "昵称": "", "邮箱": e, "密码": pwd_map.get(e, "")})

    cmd = PY + [str(POST_COMMENTS),
                "--batch", str(BATCH_OUT),
                "--tokens", str(TOKENS),
                "--accounts", str(accounts_csv),
                "--delay", str(args.delay)]
    print(f"\n[exec] {' '.join(str(c) for c in cmd)}")
    rc = subprocess.call(cmd, cwd=str(ROOT))
    return rc


if __name__ == "__main__":
    sys.exit(main())
